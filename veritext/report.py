"""DOCX report builders (python-docx). Reports carry no hosting/platform references."""
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from . import APP_NAME
from .similarity import level_for

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
LEVEL_FILL = {"High": "F8D7DA", "Moderate": "FFF3CD", "Low": "D4EDDA"}


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_cell(cell, text, bold=False, color=None, size=9.5, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold, run.font.size = bold, Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if fill:
        _shade(cell, fill)


def _widths(t, widths_cm):
    t.autofit = False
    for i, w in enumerate(widths_cm):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def _table(doc, headers, rows, widths_cm, level_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        _set_cell(t.rows[0].cells[i], h, bold=True, color=RGBColor(255, 255, 255), fill="1F3A5F")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            fill = LEVEL_FILL.get(str(v)) if i == level_col else None
            _set_cell(cells[i], v, fill=fill)
    _widths(t, widths_cm)
    return t


def _kv_table(doc, pairs, w=(4.5, 12.0)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in pairs:
        c = t.add_row().cells
        _set_cell(c[0], k, bold=True, fill="EEF2F7")
        _set_cell(c[1], v)
    _widths(t, w)
    return t


def _h(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY
        r.font.size = Pt(14)
    return p


def _para(doc, text, italic=False, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(text)
    r.italic, r.font.size = italic, Pt(size)
    return p


def _bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(it).font.size = Pt(10.5)


def _field(run, instr):
    for kind, txt in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _new_doc(meta):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    cp = d.core_properties                       # no tool/library names in file properties
    cp.author = meta.get("institution") or APP_NAME
    cp.last_modified_by = cp.author
    cp.comments = ""
    cp.title = meta.get("title", "Similarity Report")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = (meta.get("footer") or "").strip()
    r = fp.add_run(f"{footer_text}  |  Page " if footer_text else "Page ")
    r.font.size = Pt(8)
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    _field(r2, "PAGE")
    return d


def _title_block(d, meta, subtitle=None):
    inst = meta.get("institution")
    if inst:
        p = d.add_paragraph()
        r = p.add_run(inst.upper())
        r.bold, r.font.size = True, Pt(10)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    p = d.add_paragraph()
    r = p.add_run(meta.get("title", "Similarity Report").upper())
    r.bold, r.font.size = True, Pt(20)
    r.font.color.rgb = NAVY
    sub = meta.get("subtitle") if subtitle is None else subtitle
    if sub:
        _para(d, sub, italic=True)



def _lvl_reading(lvl, n_sources):
    if n_sources == 0:
        return "No matching text found in the pages checked."
    return {"Low": "Minor overlap found. Review the matched passages listed below.",
            "Moderate": "Notable overlap with online sources. Review the sources listed below.",
            "High": "Substantial overlap with online sources. Review the sources listed below."}[lvl]


def _assess(pct):
    if pct >= 15:
        return "Substantial overlap. Review"
    if pct >= 5:
        return "Partial overlap"
    return "Minor overlap"


def build_single_report(meta, doc, web, rep, settings_note, high=40.0, moderate=15.0) -> bytes:
    """Module 1: plagiarism check of one paper. doc = Doc, web = check_web result, rep = internal_repetition."""
    import io
    d = _new_doc(meta)
    pct = web["total_pct"]
    lvl = level_for(pct, high, moderate)
    srcs = web["sources"]
    _title_block(d, meta)
    _h(d, "1. Document details")
    rows = [("Document", doc.name), ("Length analysed", f"{doc.length:,} words"),
            ("Date of check", datetime.now().strftime("%d %B %Y")),
            ("Method", "Web phrase search, source-page comparison, internal repetition analysis"),
            ("Settings", settings_note)]
    if meta.get("prepared_by"):
        rows.insert(3, ("Checked by", meta["prepared_by"]))
    _kv_table(d, rows)

    _h(d, "2. Summary of results")
    longest = max([s["longest_run"] for s in srcs] or [0])
    _kv_table(d, [
        ("Confirmed matched text",
         f"{len(srcs)} source(s) with matching text; longest matching run {longest} words" if srcs
         else f"None found in the passages checked (0 matches above {meta.get('min_run', 8)} words)"),
        ("Indicative similarity", f"{pct:.1f}% from confirmed matches" if srcs else "Below 1% from confirmed matches"),
        ("Search coverage", f"{web['queries_run']} distinctive passages searched, {web['pages_checked']} pages compared"),
        ("Internal repetition", f"{rep['count']} repeated {meta.get('n', 6)}-word phrases ({rep['pct']:.1f}% of the text)"),
        ("Overall reading", _lvl_reading(lvl, len(srcs)))])
    _para(d, "Reading of the result: shared dates, names and events appear in any text on the same subject, "
             "so each match should be judged by a person.")

    _h(d, "3. Passages searched")
    prow = []
    for i, q in enumerate(web["per_query"]):
        if q["n_results"] == 0:
            res = "No page returned"
        elif q["best_pct"] <= 0:
            res = "No match; other pages only"
        else:
            res = f"Match {q['best_pct']:.1f}% in {q['best_domain']}"
        phrase = q["phrase"] if len(q["phrase"]) < 70 else q["phrase"][:67] + "..."
        prow.append((i + 1, f"{q['pos']}%", phrase, res))
    if prow:
        _table(d, ["#", "Position", "Passage tested", "Result"], prow, [0.8, 1.9, 7.7, 6.6])
    else:
        _para(d, "No passage was long enough to search.")

    _h(d, "4. Candidate sources compared")
    if srcs:
        _table(d, ["Source", "Match", "Longest run (words)", "Assessment"],
               [((s["title"][:60] + " | " + s["url"][:60]), f"{s['pct']:.1f}%", s["longest_run"], _assess(s["pct"]))
                for s in srcs[:15]], [8.0, 1.6, 2.6, 4.8])
    else:
        _para(d, "No source with matching text was found among the pages compared.")

    if srcs:
        _h(d, "5. Matched passages")
        for s in srcs[:8]:
            for ex in s["excerpts"][:2]:
                p = d.add_paragraph(style="List Bullet")
                r = p.add_run(ex[:450] + ("..." if len(ex) > 450 else ""))
                r.font.size = Pt(9.5)
                r.font.highlight_color = 7
                p.add_run(f"\nSource: {s['url'][:90]}").font.size = Pt(8)
    _h(d, "6. Internal repetition" if srcs else "5. Internal repetition")
    if rep["count"]:
        ex = "; ".join(f'"{e}" ({k} times)' for e, k in rep["examples"][:4])
        _para(d, f"A check of the document against itself found {rep['count']} repeated {meta.get('n', 6)}-word sequences, "
                 f"for example {ex}.")
    else:
        _para(d, "No repeated passages were found inside the document.")

    _h(d, "7. Final summary: similarity percentage" if srcs else "6. Final summary: similarity percentage")
    cov = "Moderate for the passages checked" if web["queries_run"] >= 8 else "Lower, as few passages were searched"
    _kv_table(d, [("Estimated similarity", f"{pct:.1f}%" if srcs else "Less than 1% (0% confirmed matched text)"),
                  ("Level", lvl),
                  ("Basis", f"{web['queries_run']} passages searched; {web['pages_checked']} pages compared in detail"),
                  ("Internal self-overlap", f"{rep['pct']:.1f}%"),
                  ("Confidence", cov),
                  ("Note", "Independent estimate using free web tools")])
    if web["errors"]:
        _para(d, "Notes: " + "; ".join(web["errors"][:3]), italic=True, size=9)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_class_report(meta, rows, pairs, settings_note, n_files, web_run=False, top_pairs_detail=8,
                       high=40.0, moderate=15.0) -> bytes:
    """Module 2: assignment similarity. rows come from analyse_class (+ combine_web when web_run)."""
    import io
    d = _new_doc(meta)
    _title_block(d, meta)
    _h(d, "1. Details")
    kv = [("Submissions compared", str(n_files)), ("Pairs compared", f"{len(pairs):,}"),
          ("Web check", "Each submission was also searched on the web" if web_run else "Not run"),
          ("Date of check", datetime.now().strftime("%d %B %Y")),
          ("Settings", settings_note),
          ("Levels", f"High: {high:.0f}% or more. Moderate: {moderate:.0f}% to {high:.0f}%. Low: below {moderate:.0f}%.")]
    if meta.get("prepared_by"):
        kv.insert(3, ("Prepared by", meta["prepared_by"]))
    _kv_table(d, kv)

    _h(d, "2. Headline")
    top = pairs[0] if pairs else None
    if top:
        _para(d, f"Highest overlap between classmates: {top.a} and {top.b} at {top.score:.1f}%.")
    n_cls = sum(1 for r in rows if r["max"] >= moderate)
    _para(d, f"{n_cls} of {n_files} students have Moderate or High similarity with a classmate.")
    if web_run:
        n_web = sum(1 for r in rows if (r.get("web_pct") or 0) >= moderate)
        _para(d, f"{n_web} of {n_files} students have Moderate or High similarity with web sources.")

    _h(d, "3. Classmate pairs with the highest similarity")
    prow = [(i + 1, p.a, p.b, f"{p.cov_a:.1f}%", f"{p.cov_b:.1f}%", p.longest_run)
            for i, p in enumerate(pairs[:25]) if p.score >= moderate] or [("-", "No pair above threshold", "", "", "", "")]
    _table(d, ["#", "Student A", "Student B", "A found in B", "B found in A", "Longest run (words)"],
           prow, [0.8, 4.2, 4.2, 2.3, 2.3, 2.8])

    _h(d, "4. Matching passages in the top pairs")
    shown = 0
    for p in pairs[:top_pairs_detail]:
        if p.score < moderate:
            break
        shown += 1
        d.add_paragraph().add_run(f"{p.a}  and  {p.b}  ({p.score:.1f}%)").bold = True
        for ex in p.excerpts[:3]:
            para = d.add_paragraph(style="List Bullet")
            r = para.add_run(ex[:400] + ("..." if len(ex) > 400 else ""))
            r.font.size = Pt(9.5)
            r.font.highlight_color = 7
    if not shown:
        _para(d, "No pair reached the Moderate level.")

    sec = 5
    if web_run:
        _h(d, f"{sec}. Text matching web sources")
        sec += 1
        any_web = False
        for r in rows:
            if (r.get("web_pct") or 0) < moderate:
                continue
            any_web = True
            d.add_paragraph().add_run(f"{r['student']}  ({r['web_pct']:.1f}% matches web sources)").bold = True
            for s in r["web_sources"][:3]:
                para = d.add_paragraph(style="List Bullet")
                para.add_run(f"{s['url'][:100]}  ({s['pct']:.1f}%)").font.size = Pt(9)
                if s["excerpts"]:
                    ex = s["excerpts"][0]
                    rr = para.add_run("\n" + ex[:300] + ("..." if len(ex) > 300 else ""))
                    rr.font.size = Pt(9)
                    rr.font.highlight_color = 7
        if not any_web:
            _para(d, "No submission reached the Moderate level against web sources.")

    _h(d, f"{sec}. Summary: students ranked by similarity")
    srows = []
    for i, r in enumerate(rows):
        overall = r.get("overall", r["max"])
        cls = f"{r['partner']} ({r['max']:.1f}%)" if r["max"] > 0 else "-"
        web = ("Not run" if r.get("web_pct") is None else f"{r['web_pct']:.1f}%") if web_run else None
        flag = r.get("flag", "Matches classmate" if r["max"] >= moderate else "-")
        srows.append([i + 1, r["student"], f"{overall:.1f}%", cls] + ([web] if web_run else []) + [flag, r["level"]])
    heads = ["#", "Student", "Highest similarity", "Closest classmate"] + (["Web match"] if web_run else []) + ["Flag", "Level"]
    widths = [0.8, 3.0, 1.9, 3.6] + ([1.5] if web_run else []) + [3.4 if not web_run else 3.0, 1.6]
    _table(d, heads, srows, widths, level_col=len(heads) - 1)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
